"""Offline color regression: python test_new_director_colors.py (no Office/model calls)."""
import json
import sys
import tempfile
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import workflow_v6_reconstruction as reconstruction
from workflow_v6_contract import new_page
from editppt.runtime.fixed_region_runtime import CONTENT_BOX, SLIDE


def check(page_plan, expected):
    with tempfile.TemporaryDirectory() as directory, ExitStack() as mocks:
        project = Path(directory)
        (project / "logo.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 20"><rect width="100" height="20"/></svg>', encoding="utf-8")
        page = new_page(1, title="Title")
        page["state"] = "accepted"
        style = {"primary_color": "#17365D", "secondary_color": "#C7352B", "background_color": "#E7F1FA", "cjk_font": "Microsoft YaHei", "title_size_pt": 28}
        state = {"pages": [page], "style_confirmation": {"contract": style}, "logo_source": {"path": "logo.svg"}}
        body = project / "body.pptx"
        deck = Presentation()
        deck.slide_width, deck.slide_height = Cm(25.4), Cm(14.288)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Cm(2), Cm(3), Cm(10), Cm(2))
        box.text = "Body"
        run = box.text_frame.paragraphs[0].runs[0]
        run.font.color.rgb, run.font.bold = RGBColor.from_string("C7352B"), True
        deck.save(body)
        authority = {"accepted_receipt": {"path": "receipt.json"}, "accepted_source_body": {"path": "source.png"}, "worker_source_body": {}, "page_plan": page_plan}
        # Scope: color/fixed-frame execution, not sealing, state persistence, or rendering.
        replacements = {
            "_load_reconstruction_state": lambda *a: state,
            "_require_final_authority": lambda *a: authority,
            "_run_post_reconstruction_visual_qa": lambda *a: {"status": "unavailable"},
            "verify_completed_page_authority": lambda *a: {"status": "verified", "authority_mode": "sealed_reconstruction", "visual_qa": {"status": "unavailable"}, "page_plan": page_plan},
            "_render_powerpoint_deck": lambda *a: {"available": False, "status": "unavailable"},
            "_officecli_validation": lambda *a: {"available": False, "status": "unavailable"},
            "load_composition_authority": lambda *a: None,
            "_structure_validation": lambda *a: {"passed": True},
            "project_emphasis_pages": lambda *a: set(),
        }
        if "primary_relationship" not in page_plan:
            def no_recalculation(*args):
                raise AssertionError("new director must not recompute emphasis pages")
            replacements["project_emphasis_pages"] = no_recalculation
        for name, value in replacements.items():
            mocks.enter_context(patch.object(reconstruction, name, value))
        reconstruction.finalize_reconstructed_page(project, page_number=1, reconstructed_body=body, commit_state=False)
        state["pages"][0]["state"] = "page_complete"
        manifest = project / "05_v6/reconstruction_runs/page_001/pages/page_001/manifest.json"
        manifest.parent.mkdir(parents=True)
        manifest.write_text(json.dumps({
            "workflow_contract_version": "fixed-canvas-cm-v2",
            "reconstruction_contract_version": "editable-image-v3",
            "slide": dict(SLIDE), "content_box": dict(CONTENT_BOX),
            "source": {"width_px": 1904, "height_px": 896},
            "text_inventory": [], "visual_inventory": [],
            "background_strategy": "native white body background",
            "quality_checks": {"font_size_calibrated": True, "visual_inventory_matched": True,
                               "background_strategy_checked": True, "shape_corner_geometry_checked": True},
            "text_boxes": [], "tables": [], "shapes": [], "images": [], "charts": [],
        }), encoding="utf-8")
        result = reconstruction.assemble_v6_deck(project)
        assert result["status"] == "validation_incomplete"
        for path in (project / "06_v6/pages/page_001/page.pptx", project / "08_candidate/deck.pptx"):
            slide = Presentation(path).slides[0]
            run = next(shape for shape in slide.shapes if shape.name == "TextBox 1").text_frame.paragraphs[0].runs[0]
            assert str(run.font.color.rgb) == expected
            assert run.font.bold is True
            assert str(slide.background.fill.fore_color.rgb) == "E7F1FA"
            assert any(shape.has_text_frame and shape.text == "Title" for shape in slide.shapes)


if __name__ == "__main__":
    check({"director_prompt": "accepted image is the body authority"}, "C7352B")
    check({"primary_relationship": {}}, "17365D")
    print("PASS: new/legacy body colors, finalize/assembly, fixed frame, no emphasis recalculation")
