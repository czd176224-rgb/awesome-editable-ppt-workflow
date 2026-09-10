"""Offline source -> proposed structure -> confirmation -> editable special pages."""
import copy
import json
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from test_confirm_ui_contract import attach_deck_plan, load_server, valid_contract
from test_workflow_v6_composition import page
from workflow_v6_composition import compose_pages, validate_structure_selection
from workflow_v6_source import initialize_v6_project
from workflow_v6_special_pages import render_special_page, SPECIAL_ROLES
from workflow_v6_reconstruction import _structure_validation
from workflow_v6_state import load


def test_structure_preserves_original_pages_and_rejects_source_loss():
    source = [page(1, "PPT页型：正文", "专题报告", "一、产业目标", "内容"),
              page(2, "PPT页型：正文", "实施措施", "二、实施路径", "内容")]
    for n, p in enumerate(source, 1):
        p["page_number"] = n
    proposal = compose_pages({"pages": source}, complete_structure=True)
    assert [p["page_role"] for p in proposal["pages"]] == ["cover", "toc", "section", "content", "section", "content", "closing"]
    assert [p["source_page_id"] for p in proposal["pages"] if p["source_page_id"] is not None] == [1, 2]
    selected = [copy.deepcopy(p) for p in proposal["pages"] if p["page_role"] != "section"]
    for n, p in enumerate(selected, 1):
        p["output_page_number"] = n
    validate_structure_selection(proposal, selected)
    for invalid in (selected[:-2] + selected[-1:], list(reversed(selected)), [{}], [None]):
        with pytest.raises(ValueError):
            validate_structure_selection(proposal, invalid)


def test_dense_opening_and_existing_chapter_are_not_replaced_or_duplicated():
    source = [page(1, "项目报告", "经营分析。" * 80),
              page(2, "PPT页型：章节", "第一章 市场分析"), page(3, "需求增长", "正文")]
    for n, p in enumerate(source, 1):
        p["page_number"] = n
    result = compose_pages({"pages": source}, complete_structure=True)
    assert result["pages"][0]["source_page_id"] is None
    assert next(p for p in result["pages"] if p["source_page_id"] == 1)["page_role"] == "content"
    assert sum(p["page_role"] == "section" for p in result["pages"]) == 1
    toc = next(p for p in result["pages"] if p["page_role"] == "toc")
    assert toc["material_source_block_ids"] == ["b-2-2"]
    plain = [page(n, "PPT页型：正文", f"内容{n}", "正文") for n in range(1, 8)]
    for n, item in enumerate(plain, 1):
        item["page_number"] = n
    fallback = compose_pages({"pages": plain}, complete_structure=True)
    assert sum(p["page_role"] == "section" for p in fallback["pages"]) == 1


def test_confirm_and_render_all_structure_roles(tmp_path: Path):
    word, logo = tmp_path / "source.docx", tmp_path / "logo.svg"
    doc = Document()
    for number, title, chapter in ((1, "产业目标", "一、产业体系"), (2, "实施措施", "二、实施计划")):
        for text in (f"PPT第{number:02d}页", "PPT页型：正文", title, chapter, "保留原页内容"):
            doc.add_paragraph(text)
    doc.save(word)
    logo.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 20"><rect width="100" height="20"/></svg>', encoding="utf-8")
    project = tmp_path / "project"
    initialize_v6_project(word, logo, project, complete_structure=True)
    server = load_server()
    client = server.create_app(str(project)).test_client()
    response = client.get("/api/recommendations")
    assert response.status_code == 200, response.get_json()
    data = response.get_json()
    selected = [{k: v for k, v in p.items() if k != "source_preview"} for p in data["composition"]["pages"]]
    payload = {**valid_contract(revision=0), "selected_director_template_id": data["recommended_template_id"],
               "director_taskbook": data["director_taskbook"], "confirmed_pages": selected, "structure_confirmed": True}
    attach_deck_plan(project, payload, structured=True)
    bad = copy.deepcopy(payload)
    bad["confirmed_pages"] = [p for p in selected if p["source_page_number"] is None]
    assert client.post("/api/confirm", json=bad).status_code == 400
    confirmed = client.post("/api/confirm", json=payload)
    assert confirmed.status_code == 200, confirmed.get_json()
    assert server._wait(project, "final", 1) == 0
    composition = json.loads((project / "02_v6/page_composition.json").read_text(encoding="utf-8"))
    assert _structure_validation(project, load(project), composition)["passed"]
    for p in composition["pages"]:
        if p["page_role"] in SPECIAL_ROLES:
            render_special_page(project, p["output_page_number"])
            deck = Presentation(project / f"06_v6/pages/page_{p['output_page_number']:03d}/page.pptx")
            assert any(s.has_text_frame for s in deck.slides[0].shapes)
    tampered = copy.deepcopy(composition)
    tampered["pages"] = tampered["pages"][:-1]
    with pytest.raises(ValueError, match="differs"):
        _structure_validation(project, load(project), tampered)
