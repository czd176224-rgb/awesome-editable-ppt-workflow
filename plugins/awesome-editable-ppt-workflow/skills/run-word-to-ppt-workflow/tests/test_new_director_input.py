"""Disposable source/UI checks. No network, model, Office or production writes."""
import json
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import deck_planning
from awesome_page_materials import collect_page_materials
from test_awesome_page_materials import _project
from workflow_v6_state import load, save


def test_original_title_is_material_and_derived_title_is_the_design_title():
    with TemporaryDirectory() as tmp:
        project = _project(Path(tmp))
        before = (project / "02_v6/paginated_word_source.json").read_bytes()
        state = load(project)
        state["pages"][0]["title"] = "全文与本页表达推导的新标题"
        save(project, state)
        result = collect_page_materials(project, 1)
        assert result["fixed_page_title"] == "全文与本页表达推导的新标题"
        assert result["complete_word_content"][0]["text"] == "第1页标题"
        assert result["complete_word_content"] == json.loads(before)["pages"][0]["blocks"]
        assert (project / "02_v6/paginated_word_source.json").read_bytes() == before


def test_full_context_and_ui_page_role_reach_the_director_boundary():
    with TemporaryDirectory() as tmp:
        project = Path(tmp)
        (project / "02_v6").mkdir()
        (project / "00_source").mkdir()
        (project / "00_source/source.docx").write_bytes(b"fixed source for mocked extraction")
        page = {"output_page_number": 1, "source_page_id": 20, "fixed_page_title": "推导标题",
                "page_role": "chapter", "chapter_title": "章节", "visible_page_number": False}
        (project / "02_v6/page_composition.json").write_text(json.dumps({"pages": [page]}), encoding="utf-8")
        full = {"pages": [{"page_number": n, "blocks": [{"text": f"完整原文{n}"}]} for n in range(1, 43)]}
        plan = {"title": "推导标题", "chapter_title": "章节"}
        with patch.object(deck_planning, "confirmed_page_plan", return_value=plan) as validate, \
             patch("extract_docx_pages.extract_auto", return_value=full) as extract:
            assert deck_planning.confirmed_page_context(project, 1) == full["pages"]
            assert deck_planning.confirmed_page_composition(project, 1) == page
            validate.reset_mock()
            extract.reset_mock()
            assert deck_planning.confirmed_page_inputs(project, 1, include_composition=True) == {
                "plan": plan, "context": full["pages"], "composition": page}
            validate.assert_called_once_with(project, 1)
            assert extract.call_count == 0  # The prior full-context read already parsed these exact bytes.
        with patch.object(deck_planning, "confirmed_page_plan", return_value={"title": "冲突标题"}), \
             patch("extract_docx_pages.extract_auto", return_value=full):
            try:
                deck_planning.confirmed_page_composition(project, 1)
            except ValueError:
                pass
            else:
                raise AssertionError("conflicting derived and outer title was accepted")
            try:
                deck_planning.confirmed_page_inputs(project, 1, include_composition=True)
            except ValueError:
                pass
            else:
                raise AssertionError("combined inputs accepted conflicting title")


def test_disposable_ui_confirmation_seals_derived_chapters_and_preserves_word():
    from docx import Document
    from codex_subscription_runtime import CodexStructuredResult
    from confirm_ui.server import create_app, _wait
    from workflow_v6_source import initialize_v6_project

    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        document = Document()
        for number in (1, 2):
            document.add_paragraph(f"第 {number} 页 PPT")
            document.add_paragraph(f"原始标题{number}")
            document.add_paragraph(f"原始事实{number}，条件未披露。")
        word, logo, project = root / "source.docx", root / "logo.svg", root / "project"
        document.save(word)
        logo.write_text('<svg xmlns="http://www.w3.org/2000/svg"/>', encoding="utf-8")
        with patch("extract_docx_pages._render_pdf_with_word", side_effect=AssertionError("offline test must not invoke Office")):
            initialize_v6_project(word, logo, project)
        before = json.loads((project / "02_v6/paginated_word_source.json").read_text(encoding="utf-8"))
        selected = json.loads((project / "02_v6/page_composition.json").read_text(encoding="utf-8"))["pages"]
        client = create_app(project).test_client()
        template = client.get("/api/recommendations").get_json()["templates"][0]
        payload = {**template["defaults"], "submission_id": "new-director-offline", "revision": 0,
                   "selected_director_template_id": template["id"], "director_taskbook": template["director_taskbook"],
                   "confirmed_pages": selected, "structure_confirmed": True}
        plan = {"pages": [{"output_page_number": n, "chapter_title": "全文理解的章节", "title": f"推导标题{n}",
                           "emphasis": "保留事实与条件", "previous_connection": "", "next_connection": ""} for n in (1, 2)]}
        def invoke(project, **kwargs):
            request = json.loads(kwargs["prompt"].split("\n", 1)[1])
            assert len(request["complete_word"]["pages"]) == 2
            assert request["selected_word"] == before
            return CodexStructuredResult(plan, "offline-thread", "offline-turn", "stub", "offline", "offline", None, {}, {})
        with patch("codex_subscription_runtime.invoke_structured", invoke), \
             patch("extract_docx_pages._render_pdf_with_word", side_effect=AssertionError("offline test must not invoke Office")):
            response = client.post("/api/deck-plan", json=payload)
        assert response.status_code == 200, response.get_json()
        payload["deck_plan"] = response.get_json()
        response = client.post("/api/confirm", json=payload)
        assert response.status_code == 200, response.get_json()
        assert _wait(project, "final", 1) == 0
        for number in (1, 2):
            page = deck_planning.confirmed_page_composition(project, number)
            materials = collect_page_materials(project, number)
            assert page["chapter_title"] == "全文理解的章节"
            assert page["fixed_page_title"] == materials["fixed_page_title"] == f"推导标题{number}"
            assert materials["complete_word_content"] == before["pages"][number - 1]["blocks"]
        import pytest
        with patch("workflow_v6_state.load", wraps=load) as read_state, \
             patch.object(deck_planning, "source_digest", wraps=deck_planning.source_digest) as hash_source, \
             patch("extract_docx_pages._render_pdf_with_word", side_effect=AssertionError("offline test must not invoke Office")):
            director_inputs = deck_planning.confirmed_page_inputs(project, 1, include_composition=True)
            assert read_state.call_count == hash_source.call_count == 1
            review_inputs = deck_planning.confirmed_page_inputs(project, 1)
            assert read_state.call_count == hash_source.call_count == 2
            assert review_inputs == {key: director_inputs[key] for key in ("plan", "context")}
            assert len(review_inputs["context"]) == 2
        source = deck_planning.source_path(project)
        original = source.read_bytes()
        source.write_bytes(original + b"changed")
        with pytest.raises(ValueError, match="校验失败"):
            deck_planning.confirmed_page_inputs(project, 1)
        source.write_bytes(original)
        state = load(project)
        state["director_confirmation"]["deck_plan"]["plan"]["pages"][0]["emphasis"] = "unsealed change"
        (project / "workflow_v6.json").write_text(json.dumps(state), encoding="utf-8")
        with pytest.raises(ValueError, match="deck plan digest is invalid"):
            deck_planning.confirmed_page_inputs(project, 1, include_composition=True)
