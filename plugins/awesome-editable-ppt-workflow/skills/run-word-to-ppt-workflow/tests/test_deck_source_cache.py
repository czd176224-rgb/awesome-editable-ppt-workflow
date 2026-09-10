"""Disposable Word inputs; no models, network or Office rendering."""
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from unittest.mock import patch

from docx import Document
import pytest

from deck_planning import complete_source
from extract_docx_pages import extract_auto
from workflow_v6_source import V6_PAGE_MARKER


def _source(project, text="Original facts", name="source.docx"):
    path = project / "00_source" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for number in (1, 2, 3):
        document.add_paragraph(f"第 {number} 页")
        document.add_paragraph(f"Title {number}")
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Condition"
    table.cell(0, 1).text = "Value"
    document.save(path)
    return path


def test_concurrent_pages_parse_once_keep_exact_source_and_return_independent_copies(tmp_path):
    source = _source(tmp_path)
    expected = extract_auto(source, marker_pattern=V6_PAGE_MARKER)
    start = Barrier(4)

    def page_request(_number):
        start.wait(timeout=10)
        return complete_source(tmp_path)

    with patch("extract_docx_pages.extract_auto", wraps=extract_auto) as parse:
        with ThreadPoolExecutor(max_workers=4) as workers:
            results = list(workers.map(page_request, range(4)))
        assert parse.call_count == 1
        assert all(result == expected for result in results)
        results[0]["pages"][0]["blocks"][0]["text"] = "Caller mutation"
        assert results[1] == complete_source(tmp_path) == expected
        assert parse.call_count == 1
        _source(tmp_path, "Changed facts")
        assert complete_source(tmp_path) == extract_auto(source, marker_pattern=V6_PAGE_MARKER)
        assert parse.call_count == 2
        source.unlink()
        with pytest.raises(FileNotFoundError):
            complete_source(tmp_path)


def test_project_identity_and_full_manuscript_selection_do_not_share_wrong_results(tmp_path):
    first, second = tmp_path / "first", tmp_path / "second"
    source = _source(first)
    other = _source(second)
    other.write_bytes(source.read_bytes())
    with patch("extract_docx_pages.extract_auto", wraps=extract_auto) as parse:
        assert complete_source(first) == complete_source(second)
        assert parse.call_count == 2
        full = _source(first, "Full manuscript", "full_deck_context.docx")
        assert complete_source(first) == extract_auto(full, marker_pattern=V6_PAGE_MARKER)
        assert parse.call_count == 3
        full.unlink()
        assert complete_source(first) == complete_source(second)
        assert parse.call_count == 3


@pytest.mark.parametrize("failure", ["parser", "source_change"])
def test_failed_or_changed_source_parse_is_not_cached(tmp_path, failure):
    _source(tmp_path)
    calls = 0

    def parse(path, **_kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            if failure == "parser":
                raise RuntimeError("parse failed")
            path.write_bytes(path.read_bytes() + b"changed during parse")
        return {"pages": [{"text": "Fresh complete source"}]}

    with patch("extract_docx_pages.extract_auto", side_effect=parse):
        error, message = ((RuntimeError, "parse failed") if failure == "parser"
                          else (ValueError, "changed during full-source parsing"))
        with pytest.raises(error, match=message):
            complete_source(tmp_path)
        assert complete_source(tmp_path) == {"pages": [{"text": "Fresh complete source"}]}
        assert calls == 2
