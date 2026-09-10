"""Disposable orchestration checks; no browser, provider or production project."""
import hashlib
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document

import workflow_v6_cli as cli
from confirm_ui import server
from test_confirm_ui_contract import attach_deck_plan


def arguments(project, *extra):
    return cli._parser().parse_args(["run", "--project", str(project), *extra])


def test_new_run_confirms_once_and_publishes_before_pipeline(tmp_path, monkeypatch):
    word, logo, project = tmp_path / "source.docx", tmp_path / "logo.svg", tmp_path / "project"
    doc = Document()
    doc.add_paragraph("第 1 页")
    doc.add_paragraph("证据支持结论")
    doc.add_paragraph("正文材料")
    doc.save(word)
    logo.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 20"><rect width="100" height="20"/></svg>')
    events = []

    def confirm(root, *_args):
        events.append("confirm")
        client = server.create_app(root).test_client()
        recommendations = client.get("/api/recommendations").get_json()
        template = next(t for t in recommendations["templates"] if t["id"] == recommendations["recommended_template_id"])
        payload = {"submission_id": "unified-run-test", "revision": recommendations["revision"],
                   **template["defaults"], "selected_director_template_id": template["id"],
                   "director_taskbook": recommendations["director_taskbook"]}
        attach_deck_plan(root, payload, structured=True)
        response = client.post("/api/confirm", json=payload)
        assert response.status_code == 200, response.get_json()
        return 0

    def pipeline(root, pages, **_kwargs):
        state = cli.load(root)
        assert state["style_confirmation"]["status"] == "confirmed"
        assert all(page["material_receipt"] for page in state["pages"])
        assert pages == [page["page_number"] for page in state["pages"]]
        events.append("pipeline")
        return SimpleNamespace(to_dict=lambda: {"assembly": {"status": "deferred"}})

    monkeypatch.setattr(server, "_start", confirm)
    monkeypatch.setattr(cli, "complete_project_real_assets", lambda *_a, **_k: events.append("assets"))
    monkeypatch.setattr(cli, "run_pages", pipeline)
    assert cli._run_project(arguments(project, "--word", str(word), "--logo", str(logo))) == 1
    assert events == ["confirm", "assets", "pipeline"]
    events.clear()
    assert cli._run_project(arguments(project)) == 1
    assert events == ["pipeline"]


def test_timeout_stops_before_materials(tmp_path, monkeypatch):
    project = tmp_path / "project"
    project.mkdir()
    (project / "workflow_v6.json").write_text("{}")
    monkeypatch.setattr(cli, "load", lambda _: {"style_confirmation": {"status": "pending"}})
    monkeypatch.setattr(server, "_start", lambda *_: 0)
    monkeypatch.setattr(server, "_wait", lambda *_: 124)
    materials = Mock(side_effect=AssertionError("must not prepare materials"))
    monkeypatch.setattr(cli, "complete_project_real_assets", materials)
    assert cli._run_project(arguments(project)) == 124
    materials.assert_not_called()


@pytest.mark.parametrize("recovery", [False, True])
def test_resume_skips_complete_pages_and_only_recovers_explicitly(tmp_path, monkeypatch, recovery):
    (tmp_path / "workflow_v6.json").write_text("{}")
    state = {"style_confirmation": {"status": "confirmed"}, "pages": [
        {"page_number": 1, "state": "page_complete", "material_receipt": {}},
        {"page_number": 2, "state": "technical_failed", "material_receipt": {}},
    ]}
    monkeypatch.setattr(cli, "load", lambda _: state)
    result = SimpleNamespace(to_dict=lambda: {"assembly": {"status": "validation_incomplete"}})
    normal, failed = Mock(return_value=result), Mock(return_value=result)
    monkeypatch.setattr(cli, "run_pages", normal)
    monkeypatch.setattr(cli, "recover_failed_pages", failed)
    extra = ("--pages", "2", "--recovery-round", "1") if recovery else ()
    assert cli._run_project(arguments(tmp_path, *extra)) == 1
    active, unused = (failed, normal) if recovery else (normal, failed)
    assert active.call_args.args[1] == [2]
    unused.assert_not_called()


@pytest.mark.parametrize("assembly,code", [
    ({"status": "complete", "release_ready": True, "final_output": {"path": "deck.pptx"}}, 0),
    ({"status": "complete", "release_ready": False, "final_output": None}, 1),
    ({"status": "deferred"}, 1),
])
def test_completed_resume_only_assembles_and_reports_actual_delivery(tmp_path, monkeypatch, assembly, code):
    import workflow_v6_reconstruction_worker as worker
    (tmp_path / "workflow_v6.json").write_text("{}")
    monkeypatch.setattr(cli, "load", lambda _: {"style_confirmation": {"status": "confirmed"}, "pages": [
        {"page_number": 1, "state": "page_complete", "material_receipt": {}},
    ]})
    monkeypatch.setattr(worker, "assemble_reconstructed_project", lambda *_: assembly)
    monkeypatch.setattr(cli, "run_pages", Mock(side_effect=AssertionError("completed page regenerated")))
    assert cli._run_project(arguments(tmp_path)) == code


def test_invalid_new_and_changed_resume_inputs_do_not_initialize(tmp_path, monkeypatch):
    initialize = Mock(side_effect=AssertionError("unexpected initialization"))
    monkeypatch.setattr(cli, "initialize_v6_project", initialize)
    for extra in [(), ("--recovery-round", "1"), ("--confirmation-timeout", "0"), ("--pages", "0")]:
        with pytest.raises(ValueError):
            cli._run_project(arguments(tmp_path, *extra))
    (tmp_path / "workflow_v6.json").write_text("{}")
    word = tmp_path / "other.docx"
    word.write_bytes(b"changed")
    monkeypatch.setattr(cli, "load", lambda _: {"word_source": {"sha256": hashlib.sha256(b"original").hexdigest()}})
    with pytest.raises(ValueError, match="locked project source"):
        cli._run_project(arguments(tmp_path, "--word", str(word)))
    initialize.assert_not_called()
